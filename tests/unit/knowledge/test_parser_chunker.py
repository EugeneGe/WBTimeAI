"""Parser / Chunker 单元测试（W5 / T2-28）。

覆盖（W5 开工前架构裁定 §2.2 / §2.5 / B8）：

- ``txt`` / ``md``：**标准库零依赖**，可解析 + 提取标题路径；
- ``html`` / ``pdf`` / ``docx``：**可选 extra** ``astra-knowledge[parsers]``，缺失时抛
  ``DomainError(DEPENDENCY_MISSING)``（**非静默**）；安装时正常解析；
- :class:`ParserRegistry` 选路 + 无解析器报错；
- :class:`FixedChunker` 分块参数（``size`` / ``overlap``）与元数据保留。
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from astra_core import DomainError, ErrorCode
from astra_knowledge.chunker import ChunkDraft, FixedChunker
from astra_knowledge.parser import (
    DocxParser,
    HtmlParser,
    MarkdownParser,
    ParserRegistry,
    PdfParser,
    TextParser,
)

FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "knowledge"


# -- txt / md（标准库） --------------------------------------------------------


def test_text_parser_reads_fixture() -> None:
    """``txt``：标准库解析，``metadata.format == "txt"``。"""
    data = (FIXTURES / "sample.txt").read_bytes()
    parsed = TextParser().parse(data, meta={})
    assert "AstraAI" in parsed.text
    assert parsed.metadata["format"] == "txt"


def test_markdown_parser_extracts_headings() -> None:
    """``md``：标准库解析 + 标题路径提取到 ``metadata.headings``。"""
    data = (FIXTURES / "sample.md").read_bytes()
    parsed = MarkdownParser().parse(data, meta={})
    assert parsed.metadata["format"] == "markdown"
    headings = parsed.metadata["headings"]
    assert "Astra Knowledge" in headings
    assert "Chunking" in headings
    assert "Notes" in headings


def test_registry_selects_by_mime() -> None:
    """注册表按 ``mime`` 选路（未安装 extra 时 txt / md 仍可用）。"""
    registry = ParserRegistry.default()
    assert registry.select(mime="text/markdown", source_type="file").name == "markdown"
    assert registry.select(mime="text/plain", source_type="file").name == "text"
    parsed = registry.parse(
        (FIXTURES / "sample.txt").read_bytes(), mime="text/plain", source_type="file"
    )
    assert "AstraAI" in parsed.text


def test_registry_raises_when_no_parser() -> None:
    """无可用解析器 → ``DomainError(VALIDATION_ERROR)``。"""
    registry = ParserRegistry.default()
    with pytest.raises(DomainError) as exc:
        registry.parse(b"\x00\x01", mime="application/octet-stream", source_type="file")
    assert exc.value.code == ErrorCode.VALIDATION_ERROR


# -- html / pdf / docx（可选 extra） ------------------------------------------


def _assert_optional(
    parser, module_name: str, mime: str, build_bytes, expected_format: str
) -> None:
    """通用：缺依赖 → ``DEPENDENCY_MISSING``；有依赖 → 正常解析。"""
    try:
        __import__(module_name)
    except ImportError:
        with pytest.raises(DomainError) as exc:
            parser.parse(b"", meta={})
        assert exc.value.code == ErrorCode.DEPENDENCY_MISSING
        assert exc.value.data.get("extra") == "parsers"
        return
    parsed = parser.parse(build_bytes(), meta={})
    assert parsed.metadata["format"] == expected_format


def test_html_parser_optional() -> None:
    """``html``：缺 ``beautifulsoup4`` → ``DEPENDENCY_MISSING``；否则解析。"""
    parser = HtmlParser()

    def _build() -> bytes:
        return b"<html><body><h1>Title</h1><p>Hello</p><script>x=1</script></body></html>"

    _assert_optional(parser, "bs4", "text/html", _build, "html")


def test_pdf_parser_optional() -> None:
    """``pdf``：缺 ``pypdf`` → ``DEPENDENCY_MISSING``；否则解析（记录页数）。"""
    parser = PdfParser()

    def _build() -> bytes:
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        buffer = io.BytesIO()
        writer.write(buffer)
        return buffer.getvalue()

    _assert_optional(parser, "pypdf", "application/pdf", _build, "pdf")


def test_docx_parser_optional() -> None:
    """``docx``：缺 ``python-docx`` → ``DEPENDENCY_MISSING``；否则解析。"""
    parser = DocxParser()

    def _build() -> bytes:
        import docx

        document = docx.Document()
        document.add_paragraph("Hello docx")
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    _assert_optional(
        parser,
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        _build,
        "docx",
    )


# -- FixedChunker -------------------------------------------------------------


def test_chunker_fixed_size_and_overlap() -> None:
    """分块：固定窗口 + 重叠；块序号连续；元数据（标题路径）保留。"""
    from astra_knowledge.parser import ParsedDocument

    text = "0123456789"  # 10 字符
    parsed = ParsedDocument(text=text, metadata={"headings": ["H1"]})
    drafts = FixedChunker(size=4, overlap=1).chunk(parsed)
    assert all(isinstance(item, ChunkDraft) for item in drafts)
    assert [item.ordinal for item in drafts] == list(range(len(drafts)))
    # step = 3；窗口 [0:4],[3:7],[6:10]
    assert drafts[0].content == "0123"
    assert drafts[1].content == "3456"
    assert drafts[2].content == "6789"
    assert drafts[0].metadata["headings"] == ["H1"]


def test_chunker_blank_text_returns_empty() -> None:
    """空白文本 → 空分块列表。"""
    from astra_knowledge.parser import ParsedDocument

    assert FixedChunker().chunk(ParsedDocument(text="   \n  ", metadata={})) == []


def test_chunker_invalid_params_raise() -> None:
    """非法参数 → ``DomainError(VALIDATION_ERROR)``。"""
    with pytest.raises(DomainError) as exc:
        FixedChunker(size=0, overlap=0)
    assert exc.value.code == ErrorCode.VALIDATION_ERROR
    with pytest.raises(DomainError):
        FixedChunker(size=4, overlap=4)


def test_chunker_override_params() -> None:
    """``chunk`` 支持按 KB 的 ``chunk_config`` 覆盖 ``size`` / ``overlap``。"""
    from astra_knowledge.parser import ParsedDocument

    parsed = ParsedDocument(text="abcdefghij", metadata={})
    drafts = FixedChunker(size=8, overlap=2).chunk(parsed, size=5, overlap=0)
    assert drafts[0].content == "abcde"
    assert drafts[1].content == "fghij"
