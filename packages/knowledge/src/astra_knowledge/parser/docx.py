"""``docx`` 解析器（可选依赖 ``python-docx``，函数级延迟导入；W5 裁定 B8）。

缺失依赖 → :class:`~astra_core.DomainError` ``DEPENDENCY_MISSING``（**非静默**）。
"""

from __future__ import annotations

import io
from typing import Any

from .base import ParsedDocument, missing_dependency_error

__all__ = ["DocxParser"]

#: 视为 Word（``.docx``）的 MIME。
_DOCX_MIMES = frozenset({"application/vnd.openxmlformats-officedocument.wordprocessingml.document"})


class DocxParser:
    """Word（``.docx``）解析器（段落文本 + 标题层级）。"""

    name = "docx"

    def supports(self, mime: str | None, source_type: str) -> bool:
        """支持 ``application/vnd.openxmlformats-officedocument.wordprocessingml.document``。"""
        return (mime or "") in _DOCX_MIMES

    def parse(self, data: bytes, *, meta: dict[str, Any] | None = None) -> ParsedDocument:
        """抽取段落文本，并收集 Heading 段落作为标题路径。

        Raises:
            DomainError: 缺少 ``python-docx``（``DEPENDENCY_MISSING``）。
        """
        try:
            import docx
        except ImportError as exc:  # pragma: no cover - 取决于是否安装 extra
            raise missing_dependency_error("python-docx", parser=self.name) from exc

        document = docx.Document(io.BytesIO(data))
        headings: list[str] = []
        texts: list[str] = []
        for paragraph in document.paragraphs:
            texts.append(paragraph.text)
            style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
            if style_name.startswith("Heading"):
                headings.append(paragraph.text)
        text = "\n".join(texts)
        metadata: dict[str, Any] = {"format": "docx", "headings": headings}
        metadata.update(meta or {})
        return ParsedDocument(text=text, metadata=metadata)
